#!/usr/bin/env python
"""Reference extractions for the bench corpus.

Usage: reference.py <file> [<file> ...]  -> writes bench/reference/<basename>.json per input.
Each JSON is {extractor_name: text, ...} plus optional numeric fields (pages, slides) or {"error": ...}.
"""
import datetime as dt
import json
import os
import subprocess
import sys
import warnings

warnings.filterwarnings("ignore")
BENCH = os.path.dirname(os.path.abspath(__file__))
REF = os.path.join(BENCH, "reference")
PANDOC_TIMEOUT = 120


def pandoc(path, fmt=None):
    cmd = ["pandoc"] + (["-f", fmt] if fmt else []) + ["-t", "plain", "--wrap=none", path]
    r = subprocess.run(cmd, capture_output=True, timeout=PANDOC_TIMEOUT)
    if r.returncode != 0:
        raise RuntimeError(r.stderr.decode("utf-8", "replace").strip()[:500])
    return r.stdout.decode("utf-8", "replace")


def fmt_cell(v):
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if (v.hour, v.minute, v.second, v.microsecond) == (0, 0, 0, 0) else v.isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, float):
        if v != v:
            return ""
        if v.is_integer() and abs(v) < 1e15:
            return str(int(v))
    try:
        import pandas as pd  # noqa
        if isinstance(v, pd.Timestamp):
            return fmt_cell(v.to_pydatetime())
        if v is pd.NaT:
            return ""
    except Exception:
        pass
    return str(v)


def ref_pdf(path):
    import pymupdf
    doc = pymupdf.open(path)
    if doc.is_encrypted and not doc.authenticate(""):
        return {"error": "encrypted"}
    texts = [doc[i].get_text() for i in range(doc.page_count)]
    return {"pymupdf": "\n".join(texts), "pages": doc.page_count}


def ref_docx(path):
    out = {}
    try:
        import docx
        from docx.oxml.ns import qn
        d = docx.Document(path)
        parts = []
        body = d.element.body
        for child in body.iterchildren():
            if child.tag == qn("w:p"):
                parts.append("".join(t.text or "" for t in child.iter(qn("w:t"))))
            elif child.tag == qn("w:tbl"):
                for tr in child.iter(qn("w:tr")):
                    cells = []
                    for tc in tr.iter(qn("w:tc")):
                        cells.append(" ".join("".join(t.text or "" for t in p.iter(qn("w:t"))) for p in tc.iter(qn("w:p"))))
                    parts.append("\t".join(cells))
        try:
            for rel in d.part.rels.values():
                if rel.reltype.endswith("/footnotes"):
                    from lxml import etree
                    root = etree.fromstring(rel.target_part.blob)
                    for fn in root.iter(qn("w:footnote")):
                        txt = "".join(t.text or "" for t in fn.iter(qn("w:t"))).strip()
                        if txt:
                            parts.append("[FOOTNOTE] " + txt)
        except Exception:
            pass
        out["python-docx"] = "\n".join(parts)
    except Exception as e:
        out["python-docx_error"] = f"{type(e).__name__}: {e}"[:500]
    try:
        out["pandoc"] = pandoc(path)
    except Exception as e:
        out["pandoc_error"] = f"{type(e).__name__}: {e}"[:500]
    return out


def ref_textutil(path):
    r = subprocess.run(["textutil", "-convert", "txt", "-stdout", path], capture_output=True, timeout=120)
    if r.returncode != 0 or (not r.stdout and r.stderr):
        return {"error": "textutil: " + r.stderr.decode("utf-8", "replace").strip()[:500]}
    return {"textutil": r.stdout.decode("utf-8", "replace")}


def ref_pptx(path):
    from pptx import Presentation
    from pptx.util import Pt  # noqa
    prs = Presentation(path)
    slides = []
    for slide in prs.slides:
        lines = []

        def walk(shapes):
            for sh in shapes:
                if sh.shape_type == 6 and hasattr(sh, "shapes"):  # group
                    walk(sh.shapes)
                    continue
                if sh.has_text_frame:
                    for p in sh.text_frame.paragraphs:
                        t = "".join(r.text for r in p.runs) or p.text
                        if t:
                            lines.append(t)
                if getattr(sh, "has_table", False) and sh.has_table:
                    for row in sh.table.rows:
                        lines.append("\t".join(c.text for c in row.cells))
        walk(slide.shapes)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            nt = slide.notes_slide.notes_text_frame.text
            if nt:
                lines.append("[NOTES] " + nt)
        slides.append("\n".join(lines))
    return {"python-pptx": "\n".join(slides), "slides": len(prs.slides)}


def ref_openpyxl(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"[SHEET {ws.title}]")
        for row in ws.iter_rows(values_only=True):
            parts.append("\t".join(fmt_cell(v) for v in row))
    return {"openpyxl": "\n".join(parts)}


def ref_pandas(path, engine):
    import pandas as pd
    sheets = pd.read_excel(path, sheet_name=None, header=None, engine=engine)
    parts = []
    for name, df in sheets.items():
        parts.append(f"[SHEET {name}]")
        for row in df.itertuples(index=False):
            parts.append("\t".join(fmt_cell(v) for v in row))
    return {f"pandas[{engine}]": "\n".join(parts)}


def ref_decoded(path):
    import chardet
    raw = open(path, "rb").read()
    boms = [(b"\xef\xbb\xbf", "utf-8-sig"), (b"\xff\xfe\x00\x00", "utf-32-le"), (b"\x00\x00\xfe\xff", "utf-32-be"), (b"\xff\xfe", "utf-16-le"), (b"\xfe\xff", "utf-16-be")]
    for bom, enc in boms:
        if raw.startswith(bom):
            return {"decoded": raw.decode(enc), "encoding": enc}
    try:
        return {"decoded": raw.decode("utf-8"), "encoding": "utf-8"}
    except UnicodeDecodeError:
        pass
    guess = chardet.detect(raw[:200000])
    enc = guess.get("encoding")
    if enc and guess.get("confidence", 0) > 0.6:
        try:
            return {"decoded": raw.decode(enc), "encoding": enc.lower(), "chardet_confidence": guess["confidence"]}
        except (UnicodeDecodeError, LookupError):
            pass
    return {"decoded": raw.decode("windows-1252", errors="replace"), "encoding": "windows-1252", "chardet_guess": enc}


def extract(path):
    ext = path.rsplit(".", 1)[-1].lower()
    if ext == "pdf":
        return ref_pdf(path)
    if ext == "docx":
        return ref_docx(path)
    if ext == "doc":
        return ref_textutil(path)
    if ext == "pptx":
        return ref_pptx(path)
    if ext == "ppt":
        return {"error": "no reference"}
    if ext in ("xlsx", "xlsm"):
        try:
            return ref_openpyxl(path)
        except Exception as e:
            return {"error": f"openpyxl: {type(e).__name__}: {e}"[:500]}
    if ext == "xls":
        return ref_pandas(path, "xlrd")
    if ext == "xlsb":
        return ref_pandas(path, "pyxlsb")
    if ext == "ods":
        return ref_pandas(path, "odf")
    if ext in ("odt", "odp"):
        try:
            return {"pandoc": pandoc(path)}
        except Exception as e:
            return {"error": "no reference" if ext == "odp" else f"pandoc: {e}"[:500]}
    if ext == "html":
        return {"pandoc": pandoc(path, "html")}
    if ext in ("csv", "json", "yaml", "yml", "md", "txt"):
        return ref_decoded(path)
    return {"error": f"unsupported extension {ext}"}


def main(paths):
    os.makedirs(REF, exist_ok=True)
    for p in paths:
        out = os.path.join(REF, os.path.basename(p) + ".json")
        try:
            res = extract(p)
        except Exception as e:
            res = {"error": f"{type(e).__name__}: {e}"[:500]}
        with open(out, "w", encoding="utf-8") as f:
            json.dump(res, f, ensure_ascii=False)


if __name__ == "__main__":
    main(sys.argv[1:])
